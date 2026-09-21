"""문서 구조를 Word 문서로 만들 때 내용과 서식이 보존되는지 확인한다."""

from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from document_helpers import ADOPT_A, choose, document_for, plan_with, review_of

from policy_signal_map.choices.models import ChoiceSet
from policy_signal_map.document.docx import render_docx, render_request_docx
from policy_signal_map.plan.models import IndicatorUse


def adopted_document():
    plan = plan_with()
    result = review_of(plan)
    choice_set = ChoiceSet()
    choose(result, choice_set, "R07", ADOPT_A, ["쿠폰 사용 실적"])
    document, _ = document_for(plan, choice_set)
    return document


def open_docx(document):
    return Document(BytesIO(render_docx(document)))


def all_text(word_document) -> str:
    paragraphs = [paragraph.text for paragraph in word_document.paragraphs]
    cells = [cell.text for table in word_document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*paragraphs, *cells])


def test_docx_has_header_sections_and_appendices():
    text = all_text(open_docx(adopted_document()))
    assert "보완 기획안 — 하반기 외국인 소비지원 쿠폰" in text
    assert "작성일 2026-09-17" in text
    assert "시연용 합성 수치 · 자료 버전 demo-001" in text
    for title in ("1. 사업 개요", "7. 성과 측정계획", "8. 추가 확인사항"):
        assert title in text
    assert "부록 1 변경 전후 대조표" in text
    assert "부록 2 데이터 분석 근거 및 한계" in text


def test_changed_paragraph_keeps_trace_and_blue_accent():
    word = open_docx(adopted_document())
    changed = next(paragraph for paragraph in word.paragraphs if "주요 지표: 쿠폰 사용 실적" in paragraph.text)
    assert "[추가]" in changed.text
    p_pr = changed._p.get_or_add_pPr()
    assert p_pr.find(qn("w:shd")).get(qn("w:fill")) == "EFF6FF"
    assert "변경 001 · 근거 DEMO-R07-SIGUNGU-강원-강릉시" in all_text(word)
    assert "R07-A-1" not in all_text(word)


def test_change_table_and_profile_are_in_docx():
    word = open_docx(adopted_document())
    text = all_text(word)
    assert len(word.tables) == 1
    assert "장 및 변경" in word.tables[0].rows[0].cells[0].text
    assert "금액·비중과 성과지표 확인" in text
    assert "강원특별자치도 강릉시 범위 참고 — 선택 지역의 진단이 아님" in text
    assert "업종 구성에서 비중이 가장 큰 항목" in text
    assert "1인당" not in text


def test_document_without_changes_says_so():
    document, _ = document_for(plan_with(indicator_use=IndicatorUse.REFERENCE))
    text = all_text(open_docx(document))
    assert "변경한 항목이 없습니다. 원안을 그대로 유지했습니다." in text
    assert "부록 3" not in text


def test_pending_is_written_as_plain_word_text():
    document, _ = document_for(plan_with(indicator_use=IndicatorUse.REFERENCE))
    text = all_text(open_docx(document))
    assert "세부 일정: 원안에 기재 없음 추가 확정 필요" in text
    assert "[추가 확정 필요]" not in text


def test_request_appendix_and_separate_request_docx():
    plan = plan_with()
    result = review_of(plan)
    choice_set = ChoiceSet()
    choose(result, choice_set, "R07", {"decision": "adopt", "option_id": "D"})
    document, _ = document_for(plan, choice_set)
    full_text = all_text(open_docx(document))
    request_text = all_text(Document(BytesIO(render_request_docx(document))))
    assert "부록 3 정밀 분석 요청서 초안" in full_text
    assert "정밀 분석 요청서 초안" in request_text
    assert "계약이나 자료 제공이 확정된 것이 아닙니다" in request_text
    assert "참고한 근거 및 한계" in request_text


def test_long_modified_sentence_is_not_truncated():
    plan = plan_with()
    result = review_of(plan)
    choice_set = ChoiceSet()
    long_text = "긴 보완 문장 " + "사업 담당자가 확인할 세부 조건을 문서에 빠짐없이 기록합니다. " * 30
    choose(
        result,
        choice_set,
        "R07",
        {"decision": "modify", "option_id": "B", "modified_text": long_text},
    )
    document, _ = document_for(plan, choice_set)
    assert long_text in all_text(open_docx(document))
