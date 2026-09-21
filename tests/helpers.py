from io import BytesIO

from docx import Document

from policy_signal_map.plan.models import PlanInput
from policy_signal_map.web.forms import parse_plan_form

VALID_FORM = {
    "action": "submit",
    "name": "하반기 외국인 소비지원 쿠폰",
    "business_type": "foreign_tourism",
    "target": "외국인 전체",
    "goals": ["foreign_share"],
    "region_level": "sigungu",
    "sido": "5100000000",
    "sigungu": "5115000000",
    "period_start": "2026-10-01",
    "period_end": "2026-12-31",
    "metrics": ["foreign_share"],
    "indicator_use": "direct",
}


def parse(form: dict) -> PlanInput:
    single = {k: v for k, v in form.items() if isinstance(v, str)}
    multi = {k: v for k, v in form.items() if isinstance(v, list)}
    return parse_plan_form(single, multi)


def docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return " ".join([*paragraphs, *cells])
