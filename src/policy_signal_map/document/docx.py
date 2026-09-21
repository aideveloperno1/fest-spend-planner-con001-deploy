"""보완 기획안 구조를 편집 가능한 Word 문서로 만든다."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .models import PENDING_MARK, PlanDocument, RequestDraft

FONT = "맑은 고딕"
BLUE = "2563EB"
BLUE_SOFT = "EFF6FF"
AMBER = RGBColor(161, 98, 7)
GRAY = RGBColor(71, 85, 105)
LINE = "D9E0E8"


def _set_run_font(run, size: float | None = None, *, bold: bool | None = None, color=None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _add_pending_text(paragraph, text: str) -> None:
    parts = text.split(PENDING_MARK)
    for index, part in enumerate(parts):
        if part:
            _set_run_font(paragraph.add_run(part), 10.5)
        if index < len(parts) - 1:
            run = paragraph.add_run("추가 확정 필요")
            _set_run_font(run, 9, bold=True, color=AMBER)


def _add_bullet(document: Document, text: str, *, changed: bool = False, added: bool = False) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.25
    if changed or added:
        p_pr = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), BLUE_SOFT)
        p_pr.append(shading)
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:color"), BLUE)
        left.set(qn("w:space"), "6")
        borders.append(left)
        p_pr.append(borders)
    _add_pending_text(paragraph, text)
    if changed or added:
        label = "변경" if changed else "추가"
        run = paragraph.add_run(f"  [{label}]")
        _set_run_font(run, 8.5, bold=True, color=RGBColor(29, 78, 216))


def _configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)


def _add_header(document: Document, plan: PlanDocument) -> None:
    title = document.add_paragraph(style="Title")
    title.add_run(plan.title)
    title.paragraph_format.space_after = Pt(10)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for index, item in enumerate(
        (f"작성일 {plan.created_on}", plan.data_notice, "상태 담당자 확인 전 초안")
    ):
        if index:
            run = meta.add_run("  |  ")
            _set_run_font(run, 9, color=RGBColor(148, 163, 184))
        run = meta.add_run(item)
        _set_run_font(run, 9, color=GRAY)

    summary = document.add_paragraph()
    summary.paragraph_format.space_before = Pt(8)
    summary.paragraph_format.space_after = Pt(14)
    for index, text in enumerate(
        (
            f"보완 반영 {len(plan.changes)}건",
            f"추가 확정 필요 {len(plan.pending)}건",
            f"원안 유지 {plan.unchanged_sections}개 장",
        )
    ):
        if index:
            summary.add_run("    ")
        run = summary.add_run(text)
        color = RGBColor(29, 78, 216) if index == 0 else AMBER if index == 1 else GRAY
        _set_run_font(run, 9.5, bold=True, color=color)


def _add_sections(document: Document, plan: PlanDocument) -> None:
    for section in plan.sections:
        document.add_heading(f"{section.number}. {section.title}", level=1)
        if not section.lines:
            _add_bullet(document, "해당 사항 없음")
            continue
        for line in section.lines:
            _add_bullet(
                document,
                line.text,
                changed=line.state == "changed",
                added=line.state == "added",
            )
            if line.state in ("changed", "added"):
                details = document.add_paragraph()
                details.paragraph_format.left_indent = Cm(0.65)
                details.paragraph_format.space_after = Pt(5)
                evidence = ", ".join(line.evidence_ids)
                value = f"{line.change_id}"
                if evidence:
                    value += f" · 근거 {evidence}"
                _set_run_font(details.add_run(value), 8.5, color=GRAY)


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), LINE)


def _fill_cell(cell, text: str, *, header=False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    _set_run_font(run, 8.5, bold=header, color=RGBColor(255, 255, 255) if header else None)
    if header:
        _set_cell_shading(cell, "334155")


def _add_changes(document: Document, plan: PlanDocument) -> None:
    document.add_page_break()
    document.add_heading("부록 1 변경 전후 대조표", level=1)
    if not plan.changes:
        document.add_paragraph("변경한 항목이 없습니다. 원안을 그대로 유지했습니다.")
        return
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    for cell, label, width in zip(table.rows[0].cells, ("장 및 변경", "변경 전", "변경 후 및 근거"), (Cm(2.6), Cm(6.2), Cm(8.2))):
        cell.width = width
        _fill_cell(cell, label, header=True)
    _set_repeat_header(table.rows[0])
    for change in plan.changes:
        cells = table.add_row().cells
        before = change.before or "추가"
        evidence = ", ".join(change.evidence_ids) or "근거 없음"
        after = (
            f"{change.after}\n\n검토: {change.rule_title}"
            f"{f' · 대안 {change.option_id}' if change.option_id else ''}"
            f"\n선택: {change.decision_label}\n근거: {evidence}"
        )
        for cell, value in zip(cells, (f"{change.section}장\n{change.change_id}", before, after)):
            _fill_cell(cell, value)


def _add_evidence(document: Document, plan: PlanDocument, heading: str = "부록 2 데이터 분석 근거 및 한계") -> None:
    document.add_heading(heading, level=1)
    profile = plan.profile_summary
    document.add_heading(f"지역 소비 프로필 {profile.region_label}", level=2)
    document.add_paragraph(
        f"자료 구분 {'합성 자료' if profile.data_kind == 'synthetic' else '실제 자료'} · 자료 버전 {profile.dataset_version}"
    )
    if profile.available:
        document.add_paragraph(f"범위 {profile.period_label} · {profile.basis_note}")
        for line in profile.lines:
            _add_bullet(document, line)
        for limitation in profile.limitations:
            _add_bullet(document, f"프로필 한계: {limitation}")
    else:
        document.add_paragraph(profile.note or "지역 소비 프로필이 없습니다.")

    for ref in plan.evidence_refs:
        document.add_heading(ref.evidence_id, level=2)
        document.add_paragraph(
            f"범위 {ref.scope_label} · 기간 {ref.period_label} · "
            f"{'합성 자료' if ref.data_kind == 'synthetic' else '실제 자료'} · 자료 버전 {ref.dataset_version}"
        )
        used = [change for change in plan.changes if ref.evidence_id in change.evidence_ids]
        for change in used:
            _add_bullet(document, f"{change.section}장 · {change.after}")
        for limitation in ref.limitations:
            _add_bullet(document, f"한계: {limitation}")
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    _set_run_font(
        note.add_run("전국 자료는 참고 범위이며 선택한 지역의 진단이 아닙니다. 금액과 비중은 서로 다른 질문에 답합니다."),
        8.5,
        color=GRAY,
    )


def _add_request_content(document: Document, draft: RequestDraft, *, appendix: bool) -> None:
    if appendix:
        document.add_heading("부록 3 정밀 분석 요청서 초안", level=1)
    document.add_paragraph("계약이나 자료 제공이 확정된 것이 아닙니다.")
    for label, value in (
        ("요청 목적", draft.purpose),
        ("분석 대상", draft.targets),
        ("대상 기간", draft.period),
        ("필요 지표", draft.metrics),
    ):
        paragraph = document.add_paragraph()
        _set_run_font(paragraph.add_run(f"{label}  "), 10.5, bold=True)
        _add_pending_text(paragraph, value)
    document.add_heading("보내기 전에 확인할 것", level=2)
    for item in draft.to_confirm:
        _add_bullet(document, item)
    document.add_paragraph("참여 점포와 참여자의 결제 실적은 현재 제공 자료로 조회할 수 없습니다.")


def render_docx(plan: PlanDocument) -> bytes:
    document = Document()
    _configure(document)
    document.core_properties.title = plan.title
    document.core_properties.subject = "보완 기획안"
    _add_header(document, plan)
    _add_sections(document, plan)
    _add_changes(document, plan)
    _add_evidence(document, plan)
    if plan.request_draft is not None:
        _add_request_content(document, plan.request_draft, appendix=True)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_request_docx(plan: PlanDocument) -> bytes:
    if plan.request_draft is None:
        raise ValueError("정밀 분석 요청서 초안이 없는 문서입니다")
    document = Document()
    _configure(document)
    document.core_properties.title = "정밀 분석 요청서 초안"
    document.add_paragraph("정밀 분석 요청서 초안", style="Title")
    _add_request_content(document, plan.request_draft, appendix=False)
    _add_evidence(document, plan, "참고한 근거 및 한계")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
